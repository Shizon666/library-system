# -*- coding: utf-8 -*-
import docx
path = r'd:\数据结构作业\图书管理系统\高级java课设论文_更新版.docx'
doc = docx.Document(path)
for i in [55, 56, 57, 109, 110, 111, 141, 142, 143, 144]:
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        t = p.text.strip()
        has_img = any(r._element.xpath('.//w:drawing') for r in p.runs)
        label = t[:70] if t else '(empty)'
        print(f'{i}: img={has_img} | {label}')
print('tables:', len(doc.tables))
img = sum(1 for r in doc.part.rels.values() if 'image' in r.reltype)
print('embedded images:', img)
